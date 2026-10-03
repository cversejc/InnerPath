from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("辰鉴_一期产品功能清单_协作维护版.docx")
PAGE_WIDTH = 6.97


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=95, bottom=85, end=95):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "D9E1E2")


def set_repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:tblHeader"))


def set_run_font(run, size=8.8, bold=False, color="263238"):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")


def set_style_font(style, size, bold=False, color="263238"):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    properties = style.element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")


def add_table(document, headers, rows, widths):
    if abs(sum(widths) - PAGE_WIDTH) > 0.01:
        raise ValueError(f"Table width {sum(widths)} does not match page width {PAGE_WIDTH}")

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    set_table_borders(table)
    set_repeat_header(table.rows[0])

    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_shading(cell, "27585A")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        set_run_font(paragraph.add_run(label), 8.8, bold=True, color="FFFFFF")

    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F3F7F6")
            cell.text = ""
            for line_index, line in enumerate(str(value).split("\n")):
                paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                set_run_font(paragraph.add_run(line), 8.45)
    return table


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), 13.2 if level == 1 else 10.6, bold=True, color="000000")
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.08
    set_run_font(paragraph.add_run(text), 8.9)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("辰鉴一期产品需求文档  |  ")
    set_run_font(run, 8, color="718080")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)
section.footer_distance = Inches(0.32)

normal = doc.styles["Normal"]
set_style_font(normal, 9.2)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.1
title_style = doc.styles["Title"]
set_style_font(title_style, 22, bold=True, color="000000")
title_style.paragraph_format.space_after = Pt(2)
title_properties = title_style.element.get_or_add_pPr()
title_border = title_properties.find(qn("w:pBdr"))
if title_border is not None:
    title_properties.remove(title_border)
for style_name, size in (("Heading 1", 13.2), ("Heading 2", 10.6)):
    style = doc.styles[style_name]
    set_style_font(style, size, bold=True, color="000000")
    style.paragraph_format.space_before = Pt(9 if style_name == "Heading 1" else 6)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.keep_with_next = True

footer = section.footer.paragraphs[0]
add_page_number(footer)

title = doc.add_paragraph(style="Title")
set_run_font(title.add_run("辰鉴一期产品需求文档"), 22, bold=True, color="000000")
title_properties = title._p.get_or_add_pPr()
title_border = title_properties.find(qn("w:pBdr"))
if title_border is not None:
    title_properties.remove(title_border)

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(6)
set_run_font(meta.add_run("开发前功能范围  |  v0.3  |  2026年10月3日"), 8.8, color="526366")

intro = doc.add_paragraph(
    "本文列出辰鉴一期希望提供的产品功能，供产品、设计、研发和运营团队共同理解产品范围。内容按使用角色与业务模块组织，聚焦用户流程和工作台能力。"
)
intro.paragraph_format.space_after = Pt(5)

add_heading(doc, "1. 产品核心流程")
add_table(
    doc,
    ["服务", "用户旅程", "交付规则"],
    [
        ["人生说明书", "完善资料与议题 → 提交申请 → 咨询师审核修正 → 最终审核 → 交付报告", "AI 仅生成内部工作初稿；用户不能自助生成或查看未审核报告。"],
        ["决策日历", "从已交付报告发起请求 → AI 生成 → 成功后自动交付", "日历不能单独生成；生成不进入人工审核流程。"],
    ],
    [1.05, 3.22, 2.70],
)

add_heading(doc, "2. 产品角色")
add_table(
    doc,
    ["角色", "主要目标", "主要功能范围"],
    [
        ["用户", "理解个人状态，并获得行动参考", "管理资料、申请报告、查看进度与交付报告、请求日历并记录行动。"],
        ["咨询师", "完成报告的专业审阅与交付", "处理报告申请、审核 AI 分析、修订内容、处理质量问题并交付。"],
        ["管理员", "保障服务顺利运营", "管理成员和权限、分配申请、查看生成任务、处理运营问题并追溯操作。"],
        ["AI 内容负责人", "维护 AI 内容生产质量", "管理工作流与 Skills、内容样例、测试运行和基础质量评估。"],
    ],
    [0.90, 2.00, 4.07],
)

add_heading(doc, "3. 用户端功能")
user_rows = [
    ["账号与个人资料", "手机号注册登录、密码找回；创建和维护称呼、性别、出生日期、历法、出生时间准确度及选填个人画像。"],
    ["首页与导航", "展示人生说明书、决策日历和个人空间的入口，并按当前资料与服务状态提供下一步入口。"],
    ["人生说明书申请", "确认出生资料，填写本次困惑、关注领域、期望获得的帮助和背景信息；支持暂存草稿与提交前确认。"],
    ["申请进度", "查看受理、处理中、待补充、已交付等状态和处理说明；在允许的状态下补充资料或撤回申请。"],
    ["报告阅读", "阅读咨询师审核后的正式报告，沿着“你是谁、卡在哪、往哪去”查看个人特质、当前议题、可能的运行模式和行动方向，并回看历史报告。"],
    ["决策日历生成", "从已交付报告发起日历请求，选择或确认生成周期、关注主题和目标；查看处理中、成功或失败状态。"],
    ["日历浏览", "按月查看整体节奏，按日查看阶段提示、摘要、适合关注的事项和行动建议。"],
    ["行动与决策记录", "围绕日期新增个人行动或决策记录，填写备注并更新进展状态；查看和维护历史记录。"],
    ["个人空间", "集中查看报告、日历、申请进度、个人档案和账号设置。"],
]
add_table(doc, ["功能模块", "用户可完成的操作"], user_rows, [1.35, 5.62])

add_heading(doc, "4. 咨询师工作台")
intro = doc.add_paragraph(
    "咨询师围绕一份报告申请完成从资料理解、专业判断到最终交付的工作。工作区采用三栏组织信息：左侧查看用户情境与依据，中间阅读和编辑分析或报告，右侧审核核心判断、风险和质量问题。"
)
intro.paragraph_format.space_after = Pt(4)
intro.paragraph_format.keep_with_next = True

add_heading(doc, "报告审核流程", level=2)
consultant_stages = [
    ["S0 资料准备", "汇总用户档案、当次问卷、背景信息和程序计算结果；标示资料来源、缺失或冲突。", "查看上下文与证据，必要时向用户请求补充。"],
    ["S1 专业基础分析", "AI 结合传统象征系统与可计算资料，形成候选判断与依据。", "确认、修改、拒绝或新增判断；核对证据与置信程度。"],
    ["S2 心理映射", "结合用户自述形成对行为倾向、体验和应对方式的候选理解。", "检查表述是否贴合用户经验，避免把解释写成诊断。"],
    ["S3 综合判断", "整合不同来源的信号，识别相互印证、差异和不确定部分。", "确认整合结论；保留无法消除的分歧，不要求 AI 强行统一。"],
    ["S4 机制与行动", "梳理可能的重复模式、卡点、资源和可尝试的行动方向。", "审核因果解释与建议，调整为具体、可选择的行动参考。"],
    ["S5 报告写作", "提供报告主题与叙事结构建议，并按章节生成内容草稿。", "选择内容重点，编辑并确认报告表达和行动计划。"],
    ["S6 质量与交付", "检查结构完整、依据可追溯、结论一致、表达安全及重复问题。", "查看阻断项和提醒；解决阻断项后完成最终审核并交付。"],
]
add_table(doc, ["阶段", "AI 与系统提供", "咨询师完成"], consultant_stages, [1.05, 2.90, 3.02])

add_heading(doc, "工作台功能", level=2)
consultant_rows = [
    ["申请队列", "查看待处理和本人负责的申请，按状态筛选，接收或打开申请；管理员可进行分配。"],
    ["申请资料与依据", "查看本次问卷、用户背景、计算结果、历史必要结论、资料来源和待核实内容。"],
    ["判断审阅", "逐项查看 AI 候选判断、证据、置信度和风险；接受、修改、拒绝或补充判断。"],
    ["AI 编辑辅助", "针对选中内容重新表达、降低确定性、补充替代解释、检查冲突与一致性、压缩或扩写。"],
    ["报告编辑", "查看建议的报告主题与章节结构，选择重点，分段生成和编辑报告内容；保留内容来源与修改记录。"],
    ["补充资料", "标明需要补充的问题并退回用户；用户补交后继续处理原申请。"],
    ["最终审核与交付", "查看结构、依据、安全和一致性检查结果；未解决的阻断项存在时不能交付。交付后报告版本只读。"],
]
add_table(doc, ["功能模块", "咨询师可完成的操作"], consultant_rows, [1.35, 5.62])

management_heading = add_heading(doc, "5. 管理运营")
management_heading.paragraph_format.page_break_before = True
admin_rows = [
    ["运营总览", "查看用户、申请、报告交付和 AI 任务运行概况，定位待处理事项。"],
    ["成员与权限", "邀请咨询师和管理员，维护后台成员角色、账号状态及访问权限。"],
    ["用户管理", "查询用户及账号状态，查看必要的服务记录，处理账号启停等运营操作。"],
    ["报告申请管理", "查看申请队列，分配咨询师，处理申请关闭或未受理情况并记录原因。"],
    ["生成任务管理", "查看报告和日历 AI 任务状态及错误信息；对符合条件的失败任务进行授权重试。"],
    ["日历运行观察", "查看日历生成结果和异常情况，协助排查运行问题；正常日历请求不进入人工审批。"],
    ["操作审计", "查询申请分配、审核、交付、任务重试和后台管理等操作记录。"],
]
add_table(doc, ["功能模块", "管理员可完成的操作"], admin_rows, [1.35, 5.62])

add_heading(doc, "6. AI 内容生产管理")
ai_rows = [
    ["工作流与 Skills", "查看和维护报告生产各阶段、Skill 指令、输入输出要求、上下文规则、安全边界与版本。已发布版本保留历史，不被新版本覆盖。"],
    ["知识与样例", "整理可供 AI 参考的知识材料和案例范例，按主题、用途和表达特点归类，并配置适用的报告环节。"],
    ["试运行与评估", "用代表性案例试运行 Skill，查看输入、输出和问题；评估依据忠实度、覆盖度、一致性、过度推断、安全表达和人工修订情况。"],
    ["问题追踪", "记录运行失败和内容质量问题，关联工作流、Skill 版本和案例，支持复核与后续评估。"],
]
add_table(doc, ["功能模块", "AI 内容负责人可完成的操作"], ai_rows, [1.35, 5.62])

add_heading(doc, "7. 产品规则与内容原则")
for item in [
    "人生说明书必须经过用户申请、咨询师审核修正和最终审核后交付；未经审核的 AI 初稿仅供内部工作使用。",
    "决策日历只能关联用户已交付的人生说明书；AI 生成成功后自动向用户交付，不经过人工审批。",
    "报告和日历提供自我理解与行动参考，不承诺必然结果，不替代用户作出现实决定。",
    "AI 生成内容应以用户提供的资料和可追溯依据为基础；专业判断由咨询师确认，资料冲突与不确定性应清楚呈现。",
    "用户只能访问自己的资料、申请、已交付报告、日历和行动记录；咨询师按申请分配范围查看处理所需信息。",
]:
    add_bullet(doc, item)

doc.core_properties.title = "辰鉴一期产品需求文档"
doc.core_properties.subject = "开发前产品功能范围"
doc.core_properties.keywords = "辰鉴, 产品需求, 功能范围"
doc.save(OUTPUT)
print(OUTPUT)
